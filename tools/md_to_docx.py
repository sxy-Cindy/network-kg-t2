#!/usr/bin/env python3
"""把课程报告 Markdown 转成排版好的 DOCX（无需 pandoc）。

    .venv/Scripts/python.exe tools/md_to_docx.py \
        report/题目二-威胁情报知识图谱实验报告.md \
        report/题目二-威胁情报知识图谱实验报告.docx

只覆盖报告实际用到的语法：标题、段落、粗体/行内代码、无序/有序列表、
引用块、围栏代码块、管道表格、水平线。A4、宋体/黑体、10 pt 正文、
表格 8~8.5 pt，页脚页码。刻意做得紧凑，保证 ≤10 页。
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HEADING_RX = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET_RX = re.compile(r"^\s*[-*]\s+(.*)$")
NUMBER_RX = re.compile(r"^\s*(\d+)\.\s+(.*)$")
TABLE_SEP_RX = re.compile(r"^\|[\s:|-]+\|$")
INLINE_RX = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
LINK_RX = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
CJK_RX = re.compile(r"[一-鿿]")


def mostly_cjk(text: str) -> bool:
    chars = [c for c in text if not c.isspace()]
    return bool(chars) and len(CJK_RX.findall(text)) / len(chars) > 0.25


def set_east_asian(run, name: str) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = Cm(1.9)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = normal.paragraph_format
    pf.line_spacing = 1.12
    pf.space_after = Pt(3)
    pf.space_before = Pt(0)

    sizes = {1: 17, 2: 13.5, 3: 11.5, 4: 10.5}
    for lvl in range(1, 5):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "黑体"
        st.font.size = Pt(sizes[lvl])
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.paragraph_format.space_before = Pt(10 if lvl <= 2 else 6)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.keep_with_next = True

    add_page_number(sec.footer.paragraphs[0])


def add_inline(par, text: str, size: float | None = None) -> None:
    text = LINK_RX.sub(r"\1", text)
    for part in INLINE_RX.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt((size or 10) - 1.2)
            run.font.color.rgb = RGBColor(0x9C, 0x27, 0x0B)
        else:
            run = par.add_run(part)
        if size:
            run.font.size = Pt(size)
        set_east_asian(run, "宋体")


def add_table(doc: Document, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not TABLE_SEP_RX.match("|" + "|".join(r) + "|")]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    cells = [r + [""] * (ncol - len(r)) for r in cells]
    table = doc.add_table(rows=len(cells), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    size = 8.0 if ncol >= 7 else 8.5
    for ri, row in enumerate(cells):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            add_inline(par, val, size=size)
            if ri == 0:
                for run in par.runs:
                    run.bold = True
                shade(cell, "E7E6E6")
        if ri == 0:
            repeat_header(table.rows[0])
        no_split(table.rows[ri])
    # 表后留一点空隙，避免与下一段贴死
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.line_spacing = 0.6


def add_code(doc: Document, lines: list[str]) -> None:
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.3)
    pf.space_before = Pt(3)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.0
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F4F4F4")
    ppr.append(shd)
    for i, line in enumerate(lines):
        run = par.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        if i < len(lines) - 1:
            run.add_break()


def add_quote(doc: Document, lines: list[str]) -> None:
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.5)
    pf.space_before = Pt(2)
    pf.space_after = Pt(4)
    parts = [l.lstrip("> ").strip() for l in lines]
    text = "".join(parts) if mostly_cjk(" ".join(parts)) else " ".join(parts)
    add_inline(par, text)
    for run in par.runs:
        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def convert(src: Path, dst: Path) -> None:
    doc = Document()
    setup(doc)
    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue
        if stripped == "---":
            i += 1
            continue

        m = HEADING_RX.match(stripped)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            if level == 1:
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.space_after = Pt(8)
                run = par.add_run(text)
                run.bold = True
                run.font.size = Pt(18)
                set_east_asian(run, "黑体")
            else:
                par = doc.add_heading(level=level - 1)
                add_inline(par, text)
                for run in par.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    set_east_asian(run, "黑体")
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code(doc, block)
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEP_RX.match(lines[i + 1].strip()):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue

        if stripped.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i])
                i += 1
            add_quote(doc, block)
            continue

        m = BULLET_RX.match(line)
        if m:
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(1)
            add_inline(par, m.group(1))
            i += 1
            continue

        m = NUMBER_RX.match(line)
        if m:
            # 不用 Word 的 List Number：它的编号跨列表连续（第二段会从 3 开始）。
            # 直接写序号并做悬挂缩进，保证每个列表都从 1 开始。
            par = doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = Cm(0.6)
            pf.first_line_indent = Cm(-0.6)
            pf.space_after = Pt(1)
            num = par.add_run(f"{m.group(1)}. ")
            num.bold = True
            add_inline(par, m.group(2))
            i += 1
            continue

        # 普通段落：合并到下一个空行或结构行
        block = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", ">", "```", "---"))
                    or BULLET_RX.match(lines[i]) or NUMBER_RX.match(lines[i])):
                break
            block.append(nxt)
            i += 1
        par = doc.add_paragraph()
        joined = "".join(block) if mostly_cjk(" ".join(block)) else " ".join(block)
        add_inline(par, joined)

    doc.core_properties.title = src.stem
    doc.core_properties.author = "知识图谱课程作业"
    doc.save(dst)
    print(f"[docx] {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
